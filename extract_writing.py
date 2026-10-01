"""
从「英语文件」里的写作资料中，为词库每个单词提取：
  - phrases       : 该单词对应的写作词组/搭配（来自王江涛核心词汇、亮点词汇等词汇表）
  - essay_sentence: 包含该单词的作文句子（来自范文、模板、方法论等 PDF/docx）
用法：
  python extract_writing.py preview   只看样例，不写库
  python extract_writing.py           写入 words.db
依赖：pip install pypdf python-docx
"""
import os
import re
import sys
import sqlite3

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, 'words.db')
SRC_DIR = os.path.join(BASE_DIR, '英语文件')

MAX_PHRASES = 5           # 每个单词最多存几个词组
MAX_SENTENCE_LEN = 200    # 作文句子最大长度

CJK = r'[\u4e00-\u9fff]'

# 句子黑名单：答题卡指令、题目要求等，不是真正的作文句子
BAD_SENTENCE = ('ANSWER SHEET', 'ANSWER SHEET 2', 'Write an essay', 'You should write',
                'Part A', 'Part B', 'Section III', 'points)', 'Directions:',
                'Write a letter', 'Suppose you are')

# 虚词：用来区分“真搭配”和“同义词罗列”
FUNC_WORDS = {'of', 'to', 'in', 'on', 'at', 'for', 'with', 'and', 'or', 'a', 'an',
              'the', 'from', 'by', 'about', 'into', 'over'}


def word_variants(word):
    """单词的有限屈折形式（避免 begin/better/authority 这类前缀误匹配）"""
    w = word.lower()
    vs = {w}
    if w.endswith('y') and len(w) > 2:
        stem = w[:-1]
        vs |= {stem + 'ies', stem + 'ied', stem + 'ier', stem + 'iest'}
    if w.endswith('e'):
        vs |= {w[:-1] + 'ing', w[:-1] + 'ed', w[:-1] + 'er'}
    else:
        vs |= {w + 'ing', w + 'ed', w + 'er'}
    vs |= {w + 's', w + 'es', w + 'ly', w + 'ment', w + 'ness'}
    return {v for v in vs if len(v) > 1}


def word_pattern(word):
    alts = '|'.join(re.escape(v) for v in sorted(word_variants(word)))
    return re.compile(r'\b(?:' + alts + r')\b', re.I)


def phrase_ok(p):
    """过滤掉「同义词罗列」这类噪声"""
    ws = p.lower().split()
    if len(ws) > 4:
        return False
    if len(ws) >= 4 and not any(w in FUNC_WORDS for w in ws):
        return False          # 4 个实词连排，多半是罗列
    return True


def sentence_ok(s):
    if not re.match(r'^[A-Z"]', s):
        return False
    if not s.rstrip().endswith(('.', '!', '?')):
        return False
    if any(b in s for b in BAD_SENTENCE):
        return False
    letters = len(re.findall(r'[A-Za-z]', s))
    if letters / len(s) < 0.7:
        return False
    return True


# ---------------- 文本读取 ----------------
def read_docx(path):
    from docx import Document
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for tb in doc.tables:
        for row in tb.rows:
            parts.append(' \t '.join(c.text for c in row.cells))
    return '\n'.join(parts)


def read_pdf(path):
    from pypdf import PdfReader
    try:
        r = PdfReader(path)
        return '\n'.join((p.extract_text() or '') for p in r.pages)
    except Exception:
        return ''


CACHE_PATH = os.path.join(BASE_DIR, '.corpus_cache.json')


def load_sources():
    """返回 [(文件名, 文本)]，带缓存，避免每次重复解析几十个文件"""
    import json
    if os.path.exists(CACHE_PATH):
        try:
            with open(CACHE_PATH, encoding='utf-8') as f:
                return json.load(f)
        except Exception:
            pass

    out = _read_sources()
    try:
        import json
        with open(CACHE_PATH, 'w', encoding='utf-8') as f:
            json.dump(out, f, ensure_ascii=False)
    except Exception:
        pass
    return out


def _read_sources():
    """返回 [(文件名, 文本)]"""
    out = []
    if not os.path.isdir(SRC_DIR):
        return out
    for name in sorted(os.listdir(SRC_DIR)):
        path = os.path.join(SRC_DIR, name)
        low = name.lower()
        if low.startswith('~$'):
            continue
        if low.endswith('.docx'):
            out.append((name, read_docx(path)))
        elif low.endswith('.pdf'):
            out.append((name, read_pdf(path)))
    return out


# ---------------- 词组解析 ----------------
def split_en_cn(text):
    """把「英文 中文」或「中文 英文」拆成 (英文, 中文)"""
    text = text.replace('\t', ' ').strip()
    en = ' '.join(re.findall(r"[A-Za-z][A-Za-z'\-]*(?:\s+[A-Za-z][A-Za-z'\-]*)*", text)).strip()
    cn = ''.join(re.findall(CJK + r"[，。、；：""''（）\s]*", text)).strip()
    return en, cn


def parse_vocab_pairs(text):
    """从词汇表文本里抽 (英文词条, 中文释义)"""
    pairs = []
    for line in text.split('\n'):
        line = line.strip()
        if not line or len(line) > 80:
            continue
        if not re.search(r'[A-Za-z]', line) or not re.search(CJK, line):
            continue
        en, cn = split_en_cn(line)
        if not en or not cn:
            continue
        if len(en.split()) > 6:          # 太长的多半不是词条
            continue
        pairs.append((en, cn))
    return pairs


def build_phrases(sources):
    """只从“词汇表类”资料里取词组"""
    keys = ('核心词汇', '亮点词汇', '同义词', '主题词')
    phrases = []
    for name, text in sources:
        if not any(k in name for k in keys):
            continue
        for en, cn in parse_vocab_pairs(text):
            phrases.append((en, cn))
    # 去重
    seen, uniq = set(), []
    for en, cn in phrases:
        k = en.lower()
        if k in seen:
            continue
        seen.add(k)
        uniq.append((en, cn))
    return uniq


# ---------------- 句子解析 ----------------
def build_sentences(sources):
    sent = []
    for name, text in sources:
        # PDF 里句子常被换行打断：段落内的单个换行先合并成空格
        text = re.sub(r'(?<!\n)\n(?!\n)', ' ', text)
        for s in re.split(r'(?<=[.!?])\s+|\n+', text):
            s = ' '.join(s.split())
            if not (40 <= len(s) <= MAX_SENTENCE_LEN):
                continue
            if re.search(CJK, s):                # 只要纯英文句子
                continue
            if not sentence_ok(s):
                continue
            sent.append(s)
    # 去重
    seen, uniq = set(), []
    for s in sent:
        k = s.lower()
        if k in seen:
            continue
        seen.add(k)
        uniq.append(s)
    return uniq


# ---------------- 匹配入库 ----------------
def collect(preview=False):
    print('正在读取写作资料（PDF/docx，约需 20~40 秒）...')
    sources = load_sources()
    usable = [(n, t) for n, t in sources if len(t) > 200]
    print(f'可读文件 {len(usable)}/{len(sources)} 个（其余为扫描版，无法提取文字）')

    phrases = build_phrases(usable)
    sentences = build_sentences(usable)
    print(f'解析到词组/词条 {len(phrases)} 条，作文句子 {len(sentences)} 句')

    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute('PRAGMA table_info(words)')
    cols = {r[1] for r in c.fetchall()}
    for col in ('phrases', 'essay_sentence'):
        if col not in cols:
            c.execute(f'ALTER TABLE words ADD COLUMN {col} TEXT')
    conn.commit()

    c.execute('SELECT id, word FROM words')
    rows = c.fetchall()

    hit_p = hit_s = 0
    samples = []
    for wid, word in rows:
        if not word or not re.match(r'^[A-Za-z]', word):
            continue
        pat = word_pattern(word)

        ps = [en for en, cn in phrases
              if phrase_ok(en) and pat.search(en) and en.lower() != word.lower()]
        ps.sort(key=lambda x: (0 if ' ' in x else 1, len(x)))     # 优先多词搭配
        ps = ps[:MAX_PHRASES]

        ss = [s for s in sentences if pat.search(s)]
        ss.sort(key=len)
        sentence = ss[0] if ss else ''

        if ps:
            hit_p += 1
        if sentence:
            hit_s += 1
        if len(samples) < 12 and (ps or sentence):
            samples.append((word, '；'.join(ps), sentence))

        if not preview:
            c.execute('UPDATE words SET phrases = ?, essay_sentence = ? WHERE id = ?',
                      ('；'.join(ps), sentence, wid))

    if not preview:
        conn.commit()
    conn.close()

    print(f'\n有词组的单词: {hit_p}/{len(rows)}')
    print(f'有作文句子的单词: {hit_s}/{len(rows)}')
    print('\n样例：')
    for w, p, s in samples:
        print(f'  【{w}】')
        if p:
            print('     词组:', p)
        if s:
            print('     句子:', s[:110])


if __name__ == '__main__':
    preview = len(sys.argv) > 1 and sys.argv[1] == 'preview'
    collect(preview=preview)
    if preview:
        print('\n（预览模式，未写入数据库；执行 python extract_writing.py 正式写入）')
