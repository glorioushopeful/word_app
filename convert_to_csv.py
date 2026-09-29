import os
import sqlite3
from openpyxl import load_workbook

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, 'words.db')
EXCEL_PATH = os.path.join(BASE_DIR, '我的词库.xlsx')


def import_from_excel():
    if not os.path.exists(EXCEL_PATH):
        print(f"❌ 找不到文件：{EXCEL_PATH}")
        return

    print("正在读取 Excel 词库...")
    wb = load_workbook(EXCEL_PATH, read_only=True)
    ws = wb.active  # 默认读取第一个工作表

    words_data = []
    # 跳过第一行表头，从第二行开始读
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] and row[1]:  # 确保单词和释义都不为空
            word = str(row[0]).strip()
            meaning = str(row[1]).strip()
            words_data.append((word, meaning))

    wb.close()
    print(f"共读取到 {len(words_data)} 个单词，开始导入数据库...")

    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()

    # 清空旧数据
    c.execute('DELETE FROM words')

    # 批量写入数据库
    c.executemany('INSERT INTO words (word, meaning, familiarity) VALUES (?, ?, 0)', words_data)
    conn.commit()
    conn.close()

    print("✅ 导入完成！你的词库已经更新。")


if __name__ == '__main__':
    import_from_excel()